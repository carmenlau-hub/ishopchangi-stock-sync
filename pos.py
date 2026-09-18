"""
pos.py — read a Mister Mobile POS stock report (stock_report*.xlsx).

Hard rules encoded here:
  * Stock figure = Column F "Available Quantity" ONLY.
    Never Column G (Quantity), never per-branch sums, never
    Quantity - Reserved - Transit. Column F already accounts for them.
  * Row 1 + row 2 are a two-level header; data starts at row 3.
  * Exclusions for iShopChangi:
        - Category = Used            (iShopChangi sells brand-new only)
        - model contains FREEBIE(S)  (giveaway units)
        - export-region sets         (JP TH TW HK CN KR MY VN US as whole word)
        - TELCO channel rows         (iShopChangi does not sell Telco sets)
        - Apple "... A" activated rows (only "... NA" is sellable)
    NOTE: unlike the other marketplaces, Apple is NOT excluded on iShopChangi.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

import openpyxl

from .brands import (
    ACTIVATED_SUFFIX,
    ACTIVATION_SPLIT_BRANDS,
    NON_ACTIVATED_SUFFIX,
    canon_brand,
)

AVAILABLE_QTY_COL = 5          # 0-based index of Column F

EXPORT_REGIONS = {"JP", "TH", "TW", "HK", "CN", "KR", "MY", "VN", "US"}

# "W US 80W CHARGER" means "bundled with a US charger" on OnePlus names.
# It must be stripped BEFORE testing for the US export-region token, but it
# also identifies a genuinely different device (see ruling 2).
US_CHARGER_RE = re.compile(r"\bW\s+US\s+\d+W\s+CHARGER\b", re.I)

MODEL_CODE_RE = re.compile(r"-[A-Z]\d{2,4}[A-Z]?$", re.I)
CAP_RAM_RE = re.compile(r"\b(\d+)\s*(GB|TB)\s*/\s*(\d+)\b", re.I)
CAP_ONLY_RE = re.compile(r"\b(\d+)\s*(GB|TB)\b", re.I)
NETWORK_RE = re.compile(r"\b(5G|4G|LTE|WIFI|WI-FI|CELL|CELLULAR|GPS|BLUETOOTH)\b", re.I)


@dataclass
class PosRow:
    stock_type_id: str
    category: str            # "New" / "Used"
    brand_raw: str
    brand: str               # canonical
    model_raw: str
    base: str                # model base, channel + model code + cap/ram stripped
    colour: str
    available: int           # Column F
    channel: str             # "PRIMARY" / "TELCO" / ""
    capacity: str            # e.g. "256GB"
    ram: str                 # e.g. "12"
    network: str             # "5G" / "4G" / "WIFI" / "CELL" / ""
    is_us_charger: bool
    excluded: str = ""       # non-empty = why it was excluded
    row_no: int = 0


@dataclass
class PosReport:
    rows: list[PosRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def usable(self) -> list[PosRow]:
        return [r for r in self.rows if not r.excluded]

    def by_id(self) -> dict[str, PosRow]:
        return {r.stock_type_id: r for r in self.rows}


def _norm_network(tok: str) -> str:
    t = (tok or "").upper().replace("-", "").replace(" ", "")
    if t in ("LTE", "4G"):
        return "4G"
    if t in ("WIFI",):
        return "WIFI"
    if t in ("CELL", "CELLULAR"):
        return "CELL"
    if t in ("GPS", "BLUETOOTH"):
        return "WIFI"          # watches: GPS/BLUETOOTH behave as the non-cell pool
    return t


def parse_model(model_raw: str) -> dict:
    """Decompose a POS model string into its match components."""
    s = " ".join(str(model_raw or "").upper().split())

    is_us_charger = bool(US_CHARGER_RE.search(s))
    s_wo_charger = US_CHARGER_RE.sub(" ", s)

    channel = ""
    for ch in ("PRIMARY", "TELCO"):
        if re.search(rf"\b{ch}\b", s_wo_charger):
            channel = ch
            s_wo_charger = re.sub(rf"\b{ch}\b", " ", s_wo_charger)

    net_m = NETWORK_RE.search(s_wo_charger)
    network = _norm_network(net_m.group(1)) if net_m else ""

    capacity = ram = ""
    m = CAP_RAM_RE.search(s_wo_charger)
    if m:
        capacity = f"{int(m.group(1))}{m.group(2).upper()}"
        ram = str(int(m.group(3)))
    else:
        m2 = CAP_ONLY_RE.search(s_wo_charger)
        if m2:
            capacity = f"{int(m2.group(1))}{m2.group(2).upper()}"

    base = s_wo_charger
    base = CAP_RAM_RE.sub(" ", base)
    base = CAP_ONLY_RE.sub(" ", base)
    base = NETWORK_RE.sub(" ", base)
    base = " ".join(base.split())
    base = MODEL_CODE_RE.sub("", base).strip()
    base = re.sub(r"\s*-\s*$", "", base).strip()

    return {
        "base": base,
        "channel": channel,
        "capacity": capacity,
        "ram": ram,
        "network": network,
        "is_us_charger": is_us_charger,
        "cleaned": s_wo_charger,
    }


def _exclusion_reason(cat: str, model_clean: str, brand: str, base: str) -> str:
    if cat.strip().lower() == "used":
        return "Used (iShopChangi sells brand-new only)"
    if re.search(r"\bFREEBIES?\b", model_clean):
        return "FREEBIE"
    for tok in re.findall(r"\b[A-Z]{2}\b", model_clean):
        if tok in EXPORT_REGIONS:
            return f"Export set ({tok})"
    if canon_brand(brand) in ACTIVATION_SPLIT_BRANDS:
        # "... NA" = non-activated (sellable). A bare trailing "A" = activated.
        parts = base.split()
        if parts and parts[-1] == ACTIVATED_SUFFIX:
            return "Activated set (iShopChangi sells non-activated only)"
        if parts and parts[-1] != NON_ACTIVATED_SUFFIX:
            # Apple accessories (AirPods, Watch) carry no NA/A marker — allow.
            pass
    return ""


def load_pos(path_or_buf) -> PosReport:
    rep = PosReport()
    wb = openpyxl.load_workbook(path_or_buf, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(rows) < 3:
        rep.errors.append("POS report has fewer than 3 rows — not a stock_report export.")
        return rep

    hdr0 = [str(c or "").strip() for c in rows[0][:6]]
    if not hdr0 or hdr0[0].lower() != "stock type id":
        rep.errors.append(
            f"POS column A header is {hdr0[0]!r}, expected 'Stock Type ID'. "
            "Is this a stock_report export?"
        )
        return rep
    sub = str(rows[1][AVAILABLE_QTY_COL] or "").strip().lower()
    if sub and "available" not in sub:
        rep.errors.append(
            f"POS Column F sub-header is {sub!r}, expected 'Available Quantity'. "
            "Column layout changed — refusing to guess."
        )
        return rep

    for i, r in enumerate(rows[2:], start=3):
        if r[0] in (None, ""):
            continue
        sid = str(r[0]).strip()
        cat = str(r[1] or "").strip()
        brand_raw = str(r[2] or "").strip()
        model_raw = str(r[3] or "").strip()
        colour = str(r[4] or "").strip()
        try:
            avail = int(float(r[AVAILABLE_QTY_COL] or 0))
        except (TypeError, ValueError):
            avail = 0
            rep.errors.append(f"POS row {i}: Column F is not numeric — treated as 0.")

        p = parse_model(model_raw)
        row = PosRow(
            stock_type_id=sid,
            category=cat,
            brand_raw=brand_raw,
            brand=canon_brand(brand_raw),
            model_raw=model_raw,
            base=p["base"],
            colour=colour,
            available=avail,
            channel=p["channel"],
            capacity=p["capacity"],
            ram=p["ram"],
            network=p["network"],
            is_us_charger=p["is_us_charger"],
            row_no=i,
        )
        row.excluded = _exclusion_reason(cat, p["cleaned"], brand_raw, p["base"])
        if not row.excluded and row.channel == "TELCO":
            row.excluded = "TELCO (iShopChangi sells PRIMARY only)"
        rep.rows.append(row)

    return rep
