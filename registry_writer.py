"""
registry_writer.py — write the updated SKU Registry workbook.

The registry is the tool's memory: re-upload it next run and every confirmed
match is already locked, so the tool never asks about it again.
"""

from __future__ import annotations

import io
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .registry import (
    SHEET_LOCKED,
    SHEET_NEW_ML,
    SHEET_NOT_SELLING,
    SHEET_NOT_YET,
    SHEET_REVIEW,
    SHEET_SUMMARY,
    SyncPlan,
    validation_summary,
)

MM_YELLOW = "FFE800"
MM_BLACK = "000000"

HDR_FILL = PatternFill("solid", fgColor=MM_BLACK)
HDR_FONT = Font(bold=True, color=MM_YELLOW, size=10)
TITLE_FONT = Font(bold=True, size=14)
BAND = PatternFill("solid", fgColor="FFFDE7")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = {
    SHEET_LOCKED: ["#", "MP Number", "Frontend Product Name (EN)", "Product UUID",
                   "SKU Code", "LOCKED Masterlist ID(s)", "ML Model(s)|Color",
                   "ML Available Qty", "Current Seller Stock", "Target Stock",
                   "# SKUs", "Status"],
    SHEET_NEW_ML: ["#", "Masterlist Stock Type ID", "Category", "Brand", "Model",
                   "Color", "Available Qty", "Suggested MP Matches",
                   "Link to MP Number", "Reviewer Decision", "Notes"],
    SHEET_REVIEW: ["#", "MP Number", "Frontend Product Name (EN)",
                   "Backend Product Name (EN)", "Product UUID", "SKU Code", "Brand",
                   "Current Seller Stock", "Suggested Masterlist Matches",
                   "Corrected Masterlist ID", "Reviewer Decision", "Notes"],
    SHEET_NOT_SELLING: ["#", "Masterlist Stock Type ID", "Category", "Brand",
                        "Model", "Color", "Available Qty"],
    SHEET_NOT_YET: ["#", "Masterlist Stock Type ID", "Category", "Brand",
                    "Model", "Color", "Available Qty"],
}

WIDTHS = {
    "#": 6, "MP Number": 14, "Frontend Product Name (EN)": 34,
    "Backend Product Name (EN)": 46, "Product UUID": 38, "SKU Code": 18,
    "LOCKED Masterlist ID(s)": 22, "ML Model(s)|Color": 62, "ML Available Qty": 15,
    "Current Seller Stock": 18, "Target Stock": 13, "# SKUs": 9, "Status": 28,
    "Masterlist Stock Type ID": 22, "Category": 12, "Brand": 16, "Model": 42,
    "Color": 22, "Available Qty": 13, "Suggested MP Matches": 46,
    "Link to MP Number": 18, "Reviewer Decision": 26, "Notes": 30,
    "Suggested Masterlist Matches": 46,
    "Corrected Masterlist ID": 22,
}


def _write_sheet(ws, headers: list[str], rows: list[dict]) -> None:
    ws.append(headers)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(h, 18)
    ws.row_dimensions[1].height = 30

    for n, r in enumerate(rows, start=2):
        ws.append([r.get(h, "") for h in headers])
        if n % 2 == 0:
            for i in range(1, len(headers) + 1):
                ws.cell(row=n, column=i).fill = BAND
        for i in range(1, len(headers) + 1):
            ws.cell(row=n, column=i).border = BORDER

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def build_registry_workbook(plan: SyncPlan, source_names: dict[str, str]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_SUMMARY

    ws["B2"] = "IShopChangi Stock Bulk Update — SKU Registry"
    ws["B2"].font = TITLE_FONT
    ws["B3"] = f"Generated {date.today():%d %b %Y}"
    ws["B3"].font = Font(italic=True, size=9, color="666666")

    r = 5
    ws[f"B{r}"], ws[f"C{r}"] = "Category", "Count"
    for cell in (ws[f"B{r}"], ws[f"C{r}"]):
        cell.fill, cell.font = HDR_FILL, HDR_FONT
    r += 1
    for k, v in validation_summary(plan).items():
        ws[f"B{r}"], ws[f"C{r}"] = k, v
        r += 1

    r += 1
    ws[f"B{r}"] = "Source files"
    ws[f"B{r}"].font = Font(bold=True)
    r += 1
    for k, v in source_names.items():
        ws[f"B{r}"], ws[f"C{r}"] = k, v
        r += 1

    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 48

    for name, rows in (
        (SHEET_LOCKED, plan.locked_rows),
        (SHEET_NEW_ML, plan.new_ml_rows),
        (SHEET_REVIEW, plan.review_rows),
        (SHEET_NOT_SELLING, plan.not_selling_rows),
        (SHEET_NOT_YET, plan.not_yet_rows),
    ):
        _write_sheet(wb.create_sheet(name), COLUMNS[name], rows)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
